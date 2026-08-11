from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/alessiofantini/Documents/Padel")
OUT = ROOT / "output" / "docx" / "CV_Alessio_Fantini_Experis_AI_Data_Integration.docx"

NAVY = "17324D"
TEAL = "2F766D"
TEXT = "243746"
MUTED = "5C6973"
LIGHT = "E5ECEF"
PALE = "F3F6F8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_paragraph_bottom_border(paragraph, color=LIGHT, size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = p_pr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        p_pr.append(pbdr)
    bottom = pbdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pbdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_hyperlink(paragraph, text, url, color=TEAL, underline=False):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    else:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "none")
        r_pr.append(u)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    set_paragraph_bottom_border(p)
    return p


def add_role(doc, title, company, dates, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(10.4)
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    company_run = p.add_run(f" | {company}")
    company_run.bold = True
    company_run.font.size = Pt(9.8)
    company_run.font.color.rgb = RGBColor.from_string(TEXT)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    meta.paragraph_format.keep_with_next = True
    rr = meta.add_run(f"{dates}  |  {location}")
    rr.italic = True
    rr.font.size = Pt(8.6)
    rr.font.color.rgb = RGBColor.from_string(MUTED)
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_bullet(doc, text, compact=False):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.48)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.2 if compact else 1.8)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(text)
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor.from_string(TEXT)
    return p


def add_project(
    doc,
    name,
    descriptor,
    stack,
    bullets,
    live=None,
    source=None,
    live_label="Demo",
    source_label="Codice",
    live_text="apri",
    source_text="GitHub",
):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(10.2)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    rr = p.add_run(f" - {descriptor}")
    rr.font.size = Pt(9.2)
    rr.font.color.rgb = RGBColor.from_string(MUTED)
    for bullet in bullets:
        add_bullet(doc, bullet, compact=True)
    sp = doc.add_paragraph()
    sp.paragraph_format.left_indent = Cm(0.26)
    sp.paragraph_format.space_after = Pt(1)
    sr = sp.add_run("Stack: ")
    sr.bold = True
    sr.font.size = Pt(8.6)
    sr.font.color.rgb = RGBColor.from_string(TEAL)
    sr2 = sp.add_run(stack)
    sr2.font.size = Pt(8.6)
    sr2.font.color.rgb = RGBColor.from_string(TEXT)
    if live or source:
        links = doc.add_paragraph()
        links.paragraph_format.left_indent = Cm(0.26)
        links.paragraph_format.space_after = Pt(1)
        links.paragraph_format.keep_with_next = False
        if live:
            lr = links.add_run(f"{live_label}: ")
            lr.bold = True
            lr.font.size = Pt(8.4)
            add_hyperlink(links, live_text, live)
        if live and source:
            sep = links.add_run("   |   ")
            sep.font.size = Pt(8.4)
            sep.font.color.rgb = RGBColor.from_string(MUTED)
        if source:
            lr = links.add_run(f"{source_label}: ")
            lr.bold = True
            lr.font.size = Pt(8.4)
            add_hyperlink(links, source_text, source)


doc = Document()
section = doc.sections[0]
# Named override to the compact_reference_guide preset: A4 is standard for Italian applications.
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.15)
section.bottom_margin = Cm(1.05)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)
section.header_distance = Cm(0.45)
section.footer_distance = Cm(0.45)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.2)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(2.8)
normal.paragraph_format.line_spacing = 1.05

for style_name in ("List Bullet",):
    st = styles[style_name]
    st.font.name = "Arial"
    st.font.size = Pt(9.0)
    st.font.color.rgb = RGBColor.from_string(TEXT)

doc.core_properties.title = "CV Alessio Fantini - AI Data Integration Specialist"
doc.core_properties.subject = "Candidatura Experis - AI Data Integration Specialist"
doc.core_properties.author = "Alessio Fantini"
doc.core_properties.keywords = "AI, data integration, REST API, SQL, n8n, reporting, dashboard, automation"

# Header masthead: compact-reference preset adapted for a modern ATS CV.
header_table = doc.add_table(rows=1, cols=2)
header_table.autofit = False
header_table.columns[0].width = Cm(14.75)
header_table.columns[1].width = Cm(3.55)
remove_table_borders(header_table)
cell = header_table.cell(0, 0)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_cell_shading(cell, NAVY)
set_cell_margins(cell, top=130, start=170, bottom=125, end=170)
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(1)
r = p.add_run("ALESSIO FANTINI")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(23)
r.font.color.rgb = RGBColor.from_string(WHITE)
p = cell.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("AI AUTOMATION & DATA INTEGRATION SPECIALIST")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(11.2)
r.font.color.rgb = RGBColor.from_string("9FE0D6")
p = cell.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Roma | Disponibile per Milano e modalità ibrida | +39 366 266 4944 | alessiofant@gmail.com")
r.font.size = Pt(8.6)
r.font.color.rgb = RGBColor.from_string(WHITE)
p = cell.add_paragraph()
p.paragraph_format.space_after = Pt(0)
add_hyperlink(p, "LinkedIn", "https://www.linkedin.com/in/alessio-fantini-3a40a6220", color="9FE0D6")
r = p.add_run("  |  ")
r.font.color.rgb = RGBColor.from_string(WHITE)
add_hyperlink(p, "GitHub", "https://github.com/Jockeys97", color="9FE0D6")
r = p.add_run("  |  ")
r.font.color.rgb = RGBColor.from_string(WHITE)
add_hyperlink(p, "Portfolio", "https://jockeys97.github.io/Jockeys97-github.io/", color="9FE0D6")

photo_cell = header_table.cell(0, 1)
photo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_cell_shading(photo_cell, NAVY)
set_cell_margins(photo_cell, top=100, start=70, bottom=100, end=120)
photo_p = photo_cell.paragraphs[0]
photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
photo_p.paragraph_format.space_after = Pt(0)
photo_p.add_run().add_picture(
    "/var/folders/r0/t9nc054j0vzdsp30zr1032r00000gn/T/codex-clipboard-d4f894ef-1da2-41c3-9650-e6ba6e613948.png",
    width=Cm(2.85),
)

add_section_heading(doc, "Profilo")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run(
    "Full-Stack & Automation Developer con esperienza concreta nell'integrazione di API REST, "
    "workflow n8n, database SQL e soluzioni AI-driven. Ho sviluppato gestionali, dashboard e "
    "automazioni per centralizzare dati, monitorare processi e supportare decisioni operative. "
    "L'esperienza in CRM e contract operations completa il profilo con analisi delle fonti, "
    "validazione delle informazioni con gli stakeholder e documentazione dei flussi. Inglese B2."
)
r.font.size = Pt(9.25)

add_section_heading(doc, "Competenze chiave")
skills_table = doc.add_table(rows=4, cols=2)
skills_table.autofit = False
skills_table.columns[0].width = Cm(4.25)
skills_table.columns[1].width = Cm(13.7)
remove_table_borders(skills_table)
skill_rows = [
    ("Data & Integration", "REST API, JSON, Webhook, SQL, integrazione database, Strapi, validazione e troubleshooting dei dati"),
    ("AI & Automation", "n8n, workflow e agenti AI, MCP, prompt engineering, knowledge base, familiarità con RAG e vector DB"),
    ("Development", "Angular, React, TypeScript, JavaScript ES6+, Node.js, PHP, Git, architetture headless"),
    ("Reporting & Operations", "Dashboard e monitoraggio, analisi requisiti, process mapping, CRM, documentazione tecnica e operativa"),
]
for idx, (label, value) in enumerate(skill_rows):
    left, right = skills_table.rows[idx].cells
    set_cell_shading(left, PALE)
    set_cell_margins(left, top=55, start=90, bottom=55, end=80)
    set_cell_margins(right, top=55, start=100, bottom=55, end=80)
    lp = left.paragraphs[0]
    lp.paragraph_format.space_after = Pt(0)
    lr = lp.add_run(label)
    lr.bold = True
    lr.font.size = Pt(8.7)
    lr.font.color.rgb = RGBColor.from_string(TEAL)
    rp = right.paragraphs[0]
    rp.paragraph_format.space_after = Pt(0)
    rr = rp.add_run(value)
    rr.font.size = Pt(8.7)
    rr.font.color.rgb = RGBColor.from_string(TEXT)

add_section_heading(doc, "Esperienza professionale")
add_role(
    doc,
    "Full Stack Developer",
    "Accademia Italiana Fitness",
    "Set 2025 - oggi",
    "Roma",
    [
        "Sviluppo e manutenzione di gestionali web con frontend Angular e API backend PHP/Node.js per dati, utenti e processi aziendali.",
        "Centralizzazione dei contenuti e dei dati tramite Strapi e architetture headless, con integrazione di database e servizi REST.",
        "Progettazione di workflow n8n per orchestrare scambi dati, webhook e logiche AI-driven dedicate ad automazione, scoring e ottimizzazione operativa.",
        "Testing, troubleshooting e monitoraggio delle integrazioni, con raccolta requisiti e verifica dei flussi insieme agli utenti interni.",
    ],
)
add_role(
    doc,
    "CRM & Contract Operations Specialist",
    "Team Service Soc. Cons. a r.l.",
    "Apr 2024 - Giu 2025",
    "Roma",
    [
        "Gestione e ottimizzazione di workflow CRM in ARXivar e Archibus per contratti, scadenze e documentazione assicurativa.",
        "Configurazione di percorsi approvativi, notifiche e dashboard di controllo per migliorare aggiornamento, tracciabilità e monitoraggio dei dati.",
        "Validazione delle informazioni e coordinamento con RUP, broker e clienti istituzionali, traducendo esigenze operative in flussi strutturati.",
    ],
)

doc.add_page_break()

add_section_heading(doc, "Progetti selezionati")
add_project(
    doc,
    "Green Active Mobile App",
    "Progetto professionale pubblicato su iOS e Android",
    "Flutter, Strapi, REST API, SQL, n8n, workflow AI",
    [
        "Contributo attivo al rinnovo dell'app mobile: funzionalità di onboarding, gestione trainer, allenamenti, monitoraggio dei progressi e integrazione con il backend Strapi.",
        "Sviluppo e verifica di flussi dati e automazioni operative, con troubleshooting multipiattaforma e rilascio su un prodotto reale utilizzato dagli utenti.",
    ],
    live="https://apps.apple.com/it/app/green-active/id1631142492",
    source="https://play.google.com/store/apps/details?id=it.greenactive.app",
    live_label="iOS",
    source_label="Android",
    live_text="Apple App Store",
    source_text="Google Play",
)
add_project(
    doc,
    "SIGNAL",
    "AI Operational Intelligence",
    "React 19, TypeScript, Express, SQLite, REST API, workflow AI",
    [
        "Dashboard operativa per analizzare eventi e segnali aziendali, interpretare le cause e proporre azioni raccomandate.",
        "Architettura predisposta per n8n, webhook e knowledge base; timeline degli eventi e interfaccia di monitoraggio delle automazioni.",
    ],
    live="https://jockeys97.github.io/SIGNAL/",
    source="https://github.com/Jockeys97/SIGNAL",
)
add_project(
    doc,
    "Steam Circuit Padel Pro",
    "Gioco web arcade 2 contro 2",
    "HTML5 Canvas, JavaScript ES6+, CSS, fisica custom e gameplay AI",
    [
        "Videogioco di padel con servizi diagonali, rimbalzi su vetro e rete, caricamento del tiro, mira, slice e cambio giocatore.",
        "Avversari e compagni controllati dall'AI, interfaccia responsive e tre arene con differenti modificatori di gameplay.",
    ],
    live="https://steam-circuit-padel-pro-jockeys97-alessios-projects-f60f895d.vercel.app",
    source="https://github.com/Jockeys97/steam-circuit-padel-pro",
)

add_section_heading(doc, "Esperienze precedenti")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Digital Marketing Assistant | Rational Tomato | Ago 2023 - Dic 2023 | Lisbona")
r.bold = True
r.font.size = Pt(9.1)
r.font.color.rgb = RGBColor.from_string(NAVY)
add_bullet(doc, "Analisi delle performance digitali, reportistica e coordinamento operativo in ambiente internazionale.", compact=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Social Media Intern | Casa Rossa | Set 2022 - Dic 2022 | Roma")
r.bold = True
r.font.size = Pt(9.1)
r.font.color.rgb = RGBColor.from_string(NAVY)

add_section_heading(doc, "Formazione")
education = [
    ("Laurea Magistrale con lode - Marketing & Digital Communication", "Università LUMSA | 2021 - 2023 | Roma"),
    ("Laurea Triennale - Scienze Politiche", "Università Roma Tre | 2016 - 2020 | Roma"),
    ("Diploma di Liceo Scientifico", "Istituto Lucio Anneo Seneca | 2012 - 2016 | Roma"),
]
for title, meta in education:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(meta)
    r.font.size = Pt(8.4)
    r.font.color.rgb = RGBColor.from_string(MUTED)

add_section_heading(doc, "Certificazioni e lingue")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Formazione tecnica: ")
r.bold = True
r.font.color.rgb = RGBColor.from_string(TEAL)
p.add_run("Complete Web Developer - Zero to Mastery; AI Prompt Engineering Bootcamp; Asfaleia Cybersecurity Course.")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Lingue: ")
r.bold = True
r.font.color.rgb = RGBColor.from_string(TEAL)
p.add_run("Italiano madrelingua; Inglese B2 (CLIC Language Center B2.1; Wall Street English Level 11).")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Disponibilità: Milano, modalità ibrida | Inserimento full-time")
r.bold = True
r.font.size = Pt(9.0)
r.font.color.rgb = RGBColor.from_string(NAVY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(0)
fr = fp.add_run("Alessio Fantini | AI Automation & Data Integration | ")
fr.font.name = "Arial"
fr.font.size = Pt(7.5)
fr.font.color.rgb = RGBColor.from_string(MUTED)
add_field(fp, "PAGE")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)

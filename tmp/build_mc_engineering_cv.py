from pathlib import Path

from docx import Document


ROOT = Path("/Users/alessiofantini/Documents/Padel")
SOURCE = ROOT / "output" / "docx" / "CV_Alessio_Fantini_Experis_AI_Data_Integration.docx"
OUTPUT = ROOT / "output" / "mc_engineering" / "Curriculum Vitae Fantini Alessio.docx"


REPLACEMENTS = {
    "AI AUTOMATION & DATA INTEGRATION SPECIALIST": "AI INTEGRATION & AUTOMATION ENGINEER",
    "Roma | Disponibile per Milano e modalità ibrida | +39 366 266 4944 | alessiofant@gmail.com":
        "Roma | Disponibile full remote | +39 366 266 4944 | alessiofant@gmail.com",
    (
        "Full-Stack & Automation Developer con esperienza concreta nell'integrazione di API REST, "
        "workflow n8n, database SQL e soluzioni AI-driven. Ho sviluppato gestionali, dashboard e "
        "automazioni per centralizzare dati, monitorare processi e supportare decisioni operative. "
        "L'esperienza in CRM e contract operations completa il profilo con analisi delle fonti, "
        "validazione delle informazioni con gli stakeholder e documentazione dei flussi. Inglese B2."
    ):
        (
            "Full-Stack & Automation Developer con esperienza concreta nell'integrazione di API REST, "
            "workflow n8n e soluzioni AI-driven all'interno di applicazioni e processi aziendali. Ho "
            "sviluppato gestionali, dashboard e automazioni collegando frontend, backend, database SQL "
            "e servizi esterni. Utilizzo Python e JavaScript per integrazioni e trasformazioni dati, "
            "con attenzione a validazione, troubleshooting e manutenibilità. Inglese B2."
        ),
    "Data & Integration": "Integration & Backend",
    "REST API, JSON, Webhook, SQL, integrazione database, Strapi, validazione e troubleshooting dei dati":
        "REST API, webhook, JSON, Python, JavaScript/TypeScript, Node.js, PHP, Strapi, SQL e SQLite",
    "AI & Automation": "AI Engineering",
    "n8n, workflow e agenti AI, MCP, prompt engineering, knowledge base, familiarità con RAG e vector DB":
        "LLM e AI agents, integrazione di API AI, n8n, prompt engineering, MCP, knowledge base; familiarità con RAG e vector DB",
    "Development": "Architecture & Delivery",
    "Angular, React, TypeScript, JavaScript ES6+, Node.js, PHP, Git, architetture headless":
        "Angular, React, Flutter, API-first e headless, integrazione applicativa, Git/GitHub, CI/CD e cloud deployment",
    "Reporting & Operations": "Quality & Operations",
    "Dashboard e monitoraggio, analisi requisiti, process mapping, CRM, documentazione tecnica e operativa":
        "Validazione input/output, fallback, logging, error handling, testing API, troubleshooting e documentazione tecnica",
    "Full Stack Developer": "Full-Stack & Automation Developer",
    "Sviluppo e manutenzione di gestionali web con frontend Angular e API backend PHP/Node.js per dati, utenti e processi aziendali.":
        "Sviluppo e manutenzione di gestionali web con Angular, API PHP/Node.js, Strapi e database SQL per utenti, contenuti e processi aziendali.",
    "Centralizzazione dei contenuti e dei dati tramite Strapi e architetture headless, con integrazione di database e servizi REST.":
        "Integrazione di applicazioni e servizi tramite REST API, webhook e architetture headless, con trasformazione e validazione dei dati tra sistemi.",
    "Progettazione di workflow n8n per orchestrare scambi dati, webhook e logiche AI-driven dedicate ad automazione, scoring e ottimizzazione operativa.":
        "Progettazione di workflow n8n e soluzioni AI-driven per orchestrare API, elaborazioni LLM, notifiche, documenti e aggiornamenti applicativi.",
    "Testing, troubleshooting e monitoraggio delle integrazioni, con raccolta requisiti e verifica dei flussi insieme agli utenti interni.":
        "Testing e troubleshooting end-to-end di workflow e API, con gestione di fallback, analisi dei log e confronto con utenti tecnici e business.",
    "Gestione e ottimizzazione di workflow CRM in ARXivar e Archibus per contratti, scadenze e documentazione assicurativa.":
        "Gestione e ottimizzazione di workflow enterprise in ARXivar e Archibus per contratti, scadenze e documentazione assicurativa.",
    "Configurazione di percorsi approvativi, notifiche e dashboard di controllo per migliorare aggiornamento, tracciabilità e monitoraggio dei dati.":
        "Configurazione di percorsi approvativi, notifiche e dashboard per migliorare qualità, tracciabilità e monitoraggio dei dati tra sistemi.",
    "Validazione delle informazioni e coordinamento con RUP, broker e clienti istituzionali, traducendo esigenze operative in flussi strutturati.":
        "Analisi dei requisiti e coordinamento con RUP, broker e clienti istituzionali, traducendo esigenze operative in processi strutturati e verificabili.",
    " - Progetto professionale pubblicato su iOS e Android": " - Applicazione professionale integrata con servizi backend",
    "Flutter, Strapi, REST API, SQL, n8n, workflow AI": "Flutter, Strapi, REST API, SQL, n8n, integrazioni e workflow AI",
    "Contributo attivo al rinnovo dell'app mobile: funzionalità di onboarding, gestione trainer, allenamenti, monitoraggio dei progressi e integrazione con il backend Strapi.":
        "Contributo al rinnovo dell'app mobile con onboarding, gestione trainer e allenamenti, monitoraggio dei progressi e integrazione con API Strapi.",
    "Sviluppo e verifica di flussi dati e automazioni operative, con troubleshooting multipiattaforma e rilascio su un prodotto reale utilizzato dagli utenti.":
        "Sviluppo e verifica di flussi dati e automazioni collegate al prodotto, con troubleshooting multipiattaforma e rilascio su iOS e Android.",
    "Dashboard operativa per analizzare eventi e segnali aziendali, interpretare le cause e proporre azioni raccomandate.":
        "Applicazione AI per analizzare segnali aziendali, interpretare possibili cause e proporre azioni operative tramite un'interfaccia di monitoraggio.",
    "Architettura predisposta per n8n, webhook e knowledge base; timeline degli eventi e interfaccia di monitoraggio delle automazioni.":
        "Layer di analisi predisposto per integrazione LLM, n8n, webhook e knowledge base, con API backend, timeline eventi e workflow monitorabili.",
    "React 19, TypeScript, Express, SQLite, REST API, workflow AI":
        "React 19, TypeScript, Express, SQLite, REST API, integrazione LLM e workflow AI",
    " - Gioco web arcade 2 contro 2": " - Applicazione web con logiche AI e motore custom",
    "HTML5 Canvas, JavaScript ES6+, CSS, fisica custom e gameplay AI":
        "HTML5 Canvas, JavaScript ES6+, CSS, architettura modulare, fisica custom e gameplay AI",
    "Videogioco di padel con servizi diagonali, rimbalzi su vetro e rete, caricamento del tiro, mira, slice e cambio giocatore.":
        "Applicazione interattiva con motore fisico custom, gestione dello stato, controlli responsive e logiche di gioco modulari.",
    "Avversari e compagni controllati dall'AI, interfaccia responsive e tre arene con differenti modificatori di gameplay.":
        "Comportamenti AI per avversari e compagni, gestione degli errori di stato e distribuzione continua tramite Vercel.",
    "Disponibilità: Milano, modalità ibrida | Inserimento full-time":
        "Disponibilità: full remote | Inserimento full-time",
    "Alessio Fantini | AI Automation & Data Integration | ":
        "Alessio Fantini | AI Integration & Automation | ",
}


def iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


doc = Document(SOURCE)
replaced = set()

containers = [doc]
for section in doc.sections:
    containers.extend([section.header, section.footer])

for container in containers:
    for paragraph in iter_paragraphs(container):
        for run in paragraph.runs:
            if run.text in REPLACEMENTS:
                replaced.add(run.text)
                run.text = REPLACEMENTS[run.text]

missing = sorted(set(REPLACEMENTS) - replaced)
if missing:
    raise RuntimeError(f"Unmatched source text ({len(missing)}): {missing}")

doc.core_properties.title = "Curriculum Vitae Fantini Alessio - AI Integration Engineer"
doc.core_properties.subject = "Candidatura MC Engineering - AI Integration Engineer"
doc.core_properties.author = "Alessio Fantini"
doc.core_properties.keywords = (
    "AI Integration Engineer, LLM, AI agents, REST API, Python, n8n, SQL, microservices, cloud, CI/CD, troubleshooting"
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

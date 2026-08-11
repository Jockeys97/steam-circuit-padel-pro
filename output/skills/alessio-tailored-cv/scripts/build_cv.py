#!/usr/bin/env python3
import argparse
import json
from pathlib import Path

from docx import Document


BASE = Path(__file__).resolve().parent.parent / "assets" / "base-cv.docx"


def require_length(name, values, expected):
    if not isinstance(values, list) or len(values) != expected:
        raise ValueError(f"{name} must contain exactly {expected} items")


def replace_run_text(doc, replacements):
    found = set()

    def paragraphs(parent):
        for paragraph in parent.paragraphs:
            yield paragraph
        for table in parent.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)

    containers = [doc]
    for section in doc.sections:
        containers.extend([section.header, section.footer])

    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                if run.text in replacements:
                    found.add(run.text)
                    run.text = replacements[run.text]

    missing = sorted(set(replacements) - found)
    if missing:
        raise RuntimeError(f"Template mismatch; unmatched text: {missing}")


def remove_empty_placeholder_bullets(doc):
    for paragraph in list(doc.paragraphs):
        if paragraph.text == "[[REMOVE_EMPTY_BULLET]]":
            paragraph._element.getparent().remove(paragraph._element)


def build(spec, output):
    require_length("skills", spec["skills"], 4)
    require_length("experiences", spec["experiences"], 2)
    require_length("experiences[0].bullets", spec["experiences"][0]["bullets"], 4)
    require_length("experiences[1].bullets", spec["experiences"][1]["bullets"], 3)
    require_length("projects", spec["projects"], 3)
    for index, project in enumerate(spec["projects"]):
        require_length(f"projects[{index}].bullets", project["bullets"], 2)
    require_length("additional_experience", spec["additional_experience"], 2)
    require_length("education", spec["education"], 3)

    if len(spec["profile"].split()) > 80:
        raise ValueError("profile exceeds 80 words")

    doc = Document(BASE)
    s = spec
    ex1, ex2 = s["experiences"]
    p1, p2, p3 = s["projects"]
    a1, a2 = s["additional_experience"]
    e1, e2, e3 = s["education"]
    titles = s["section_titles"]

    replacements = {
        "IT AUTOMATION & WORKFLOW SPECIALIST": s["headline"],
        "Rome, Italy | Open to remote work | +39 366 266 4944 | alessiofant@gmail.com": s["contact_line"],
        "PROFILE": titles["profile"],
        "CORE SKILLS": titles["skills"],
        "PROFESSIONAL EXPERIENCE": titles["experience"],
        "SELECTED PROJECTS": titles["projects"],
        "ADDITIONAL EXPERIENCE": titles["additional"],
        "EDUCATION": titles["education"],
        "TRAINING AND LANGUAGES": titles["training"],
        (
            "IT Automation and Integration specialist with hands-on experience building and maintaining n8n workflows, "
            "REST API and webhook integrations, and AI-assisted operational processes. I turn manual workflows into "
            "structured automations with validation, monitoring, troubleshooting and clear documentation. My background "
            "combines application development with CRM and contract operations, giving me a strong service mindset and "
            "experience coordinating sensitive, business-critical information with internal and external stakeholders."
        ): s["profile"],
        "Automation & Integration": s["skills"][0]["label"],
        "n8n, REST APIs, webhooks, JSON, event-driven workflows, authentication concepts, Strapi and SQL": s["skills"][0]["value"],
        "Reliability & Support": s["skills"][1]["label"],
        "Workflow testing, logs and API responses, error handling, monitoring, validation, troubleshooting and user support": s["skills"][1]["value"],
        "Development": s["skills"][2]["label"],
        "JavaScript ES6+, TypeScript, Node.js, PHP, Angular, React, Flutter, Git and GitHub": s["skills"][2]["value"],
        "Operations & Documentation": s["skills"][3]["label"],
        "CRM workflows, process mapping, technical documentation, data quality and stakeholder coordination": s["skills"][3]["value"],
        "Full-Stack & Automation Developer": ex1["title"],
        " | Accademia Italiana Fitness": f" | {ex1['company']}",
        "Sep 2025 - Present  |  Rome, Italy": f"{ex1['dates']}  |  {ex1['location']}",
        "Design and maintain n8n workflows for CRM processes, customer onboarding, notifications, cross-service data synchronisation and webhook/API integrations.": ex1["bullets"][0],
        "Develop and support Angular applications and PHP/Node.js APIs connected to Strapi and SQL data, including user roles and business-process flows.": ex1["bullets"][1],
        "Implement AI-assisted workflows for chatbot operations, request classification, dynamic routing and scoring, with validation and human review where required.": ex1["bullets"][2],
        "Diagnose integration issues through workflow executions, logs and API responses; test changes, monitor outcomes and document operational behaviour.": ex1["bullets"][3],
        "CRM & Contract Operations Specialist": ex2["title"],
        " | Team Service Soc. Cons. a r.l.": f" | {ex2['company']}",
        "Apr 2024 - Jun 2025  |  Rome, Italy": f"{ex2['dates']}  |  {ex2['location']}",
        "Managed daily CRM-based operations in ARXivar and Archibus for contracts, deadlines, insurance records and structured document access.": ex2["bullets"][0],
        "Configured approval paths, notifications and monitoring dashboards to improve data accuracy, traceability and timely operational follow-up.": ex2["bullets"][1],
        "Supported users and coordinated with public-sector stakeholders, brokers and institutional clients while handling compliance-sensitive information.": ex2["bullets"][2],
        "Green Active Mobile App": p1["name"],
        " - Professional product published on iOS and Android": f" - {p1['descriptor']}",
        "Contributed to the mobile app renewal across onboarding, trainer management, workouts, progress tracking and Strapi backend integration.": p1["bullets"][0],
        "Built and validated related data flows and operational automations, with cross-platform troubleshooting on a live user-facing product.": p1["bullets"][1],
        "Flutter, Strapi, REST APIs, SQL, n8n and AI-assisted workflows": p1["stack"],
        "SIGNAL": p2["name"],
        " - AI Operational Intelligence": f" - {p2['descriptor']}",
        "Operational dashboard for analysing business events and signals, interpreting likely causes and presenting recommended actions.": p2["bullets"][0],
        "Architecture designed for n8n, webhooks and knowledge-base integrations, with event timelines and workflow monitoring interfaces.": p2["bullets"][1],
        "React 19, TypeScript, Express, SQLite, REST APIs and AI workflows": p2["stack"],
        "Steam Circuit Padel Pro": p3["name"],
        " - 2-vs-2 browser arcade game": f" - {p3['descriptor']}",
        "Implemented diagonal serves, wall and net bounces, charged shots, aiming, slice controls and player switching.": p3["bullets"][0],
        "Added AI-controlled opponents and teammates, responsive controls and three arenas with distinct gameplay modifiers.": p3["bullets"][1],
        "HTML5 Canvas, JavaScript ES6+, CSS, custom game physics and gameplay AI": p3["stack"],
        "Junior Web Developer | Zero to Mastery | Dec 2023 - Mar 2024 | Rome": a1["header"],
        "Built and maintained full-stack training projects with React, Node.js, REST APIs and reusable responsive components.": a1["bullet"] or "[[REMOVE_EMPTY_BULLET]]",
        "Digital Marketing Assistant | Rational Tomato | Aug 2023 - Dec 2023 | Lisbon": a2["header"],
        "Worked independently in an international environment, coordinating digital operations and performance reporting.": a2["bullet"] or "[[REMOVE_EMPTY_BULLET]]",
        "Master's Degree with honours - Marketing & Digital Communication": e1["title"],
        "LUMSA University | 2021 - 2023 | Rome": e1["meta"],
        "Bachelor's Degree - Political Science": e2["title"],
        "Roma Tre University | 2016 - 2020 | Rome": e2["meta"],
        "Scientific High School Diploma": e3["title"],
        "Istituto Lucio Anneo Seneca | 2012 - 2016 | Rome": e3["meta"],
        "Technical training: ": s["training_label"],
        "Complete Web Developer - Zero to Mastery; AI Prompt Engineering Bootcamp; Asfaleia Cybersecurity Course.": s["training"],
        "Languages: ": s["languages_label"],
        "Italian: native; English: B2 (CLIC Language Center B2.1; Wall Street English Level 11).": s["languages"],
        "Availability: full-time | Remote-first environment | Eligible to work in Italy": s["availability"],
        "Alessio Fantini | IT Automation & Workflow Specialist | ": s["footer_label"],
    }

    replace_run_text(doc, replacements)
    remove_empty_placeholder_bullets(doc)
    doc.core_properties.title = s["document_title"]
    doc.core_properties.subject = s["document_subject"]
    doc.core_properties.author = "Alessio Fantini"
    doc.core_properties.keywords = s["keywords"]

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description="Build Alessio Fantini's tailored CV from the retained DOCX template.")
    parser.add_argument("--spec", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    build(spec, args.out)
    print(args.out)


if __name__ == "__main__":
    main()
